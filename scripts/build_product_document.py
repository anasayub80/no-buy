"""Build the readable Word product requirements from the authoritative Markdown.

Run with the Codex primary Python runtime (python-docx required).
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/App-product-requirements.md'
OUTPUT = ROOT / 'docs/App-product-requirements.docx'

def plain(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')

def cell_style(cell, fill, white=False):
    tcpr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), fill); tcpr.append(shading)
    margins = OxmlElement('w:tcMar')
    for side in ['top', 'left', 'bottom', 'right']:
        e = OxmlElement('w:' + side); e.set(qn('w:w'), '105'); e.set(qn('w:type'), 'dxa'); margins.append(e)
    tcpr.append(margins)
    borders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
    tcpr.append(borders)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        for run in p.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string('FFFFFF' if white else '000000')

def make_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [1.5, 3.2, 2.3] if len(rows[0]) == 3 else [2.0, 5.0]
    if rows[0][0] == 'Area': widths = [1.35, 2.5, 3.15]
    for i, values in enumerate(rows):
        row = table.rows[i]
        props = row._tr.get_or_add_trPr()
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); props.append(repeat)
        no_split = OxmlElement('w:cantSplit'); props.append(no_split)
        for j, value in enumerate(values):
            cell = row.cells[j]; cell.width = Inches(widths[j]); cell.text = plain(value)
            cell_style(cell, '17324D' if i == 0 else ('F2F5F7' if i % 2 else 'FFFFFF'), i == 0)
            if i == 0:
                for run in cell.paragraphs[0].runs: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.7); section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.75); section.right_margin = Inches(.75)
    section.footer_distance = Inches(.3)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [('Title', 24), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]; style.font.name = 'Arial'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != 'Title'; style.font.underline = False
        style.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7 if name != 'Title' else 14)
        style.paragraph_format.keep_with_next = True
    doc.core_properties.title = 'Lessly App Idea and Product Requirements'
    doc.core_properties.subject = 'Community-informed mobile app product and implementation specification'
    doc.core_properties.author = 'Anas Ayub'
    lines = SOURCE.read_text().splitlines(); i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                values = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in values): rows.append(values)
                i += 1
            make_table(doc, rows); continue
        if line.startswith('# '): doc.add_paragraph(plain(line[2:]), 'Title')
        elif line.startswith('## '): doc.add_paragraph(plain(line[3:]), 'Heading 1')
        elif line.startswith('### '): doc.add_paragraph(plain(line[4:]), 'Heading 2')
        elif line.startswith('- '): doc.add_paragraph(plain(line[2:]), 'List Bullet')
        else: doc.add_paragraph(plain(line))
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Lessly product requirements 1.5  |  ')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    for run in footer.runs: run.font.size = Pt(8); run.font.color.rgb = RGBColor(0, 0, 0)
    # Remove inherited Word title paragraph borders, including the built-in blue rule.
    for root in (doc.styles.element, doc.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__': build()
